import json
import tempfile
from pathlib import Path

from pptmaker.m3.builder import build_pptx
from pptmaker.m3.compiler import compile_specs, normalize_layout
from pptmaker.m3.design import PRESETS
from pptmaker.m3.models import (
    ContentPlan, DesignPlan, SlideContent, SlideSpec,
)
from pptmaker.evaluation.geometry import Box, out_of_bounds, overlap_ratio


def _design() -> DesignPlan:
    return PRESETS["business"].model_copy(update={"project_id": "eval"})


def _element_box(el) -> Box:
    return Box(int(el.x * 914400), int(el.y * 914400),
               int(el.width * 914400), int(el.height * 914400))


def evaluate(cases: list[dict]) -> dict:
    specs: list[SlideSpec] = []
    layouts = set()
    crashes = []
    for case in cases:
        pages = [
            SlideContent(
                page_id=f"{case['id']}-p{i}",
                title=p["title"],
                page_type=p.get("page_type", "content"),
                layout=p.get("layout", "text"),
                body_paragraphs=p.get("body", ["text"]),
                bullet_points=p.get("bullets", []),
                image_suggestion=p.get("image_suggestion"),
                table_rows=p.get("table_rows"),
                source_references=[],
            )
            for i, p in enumerate(case["pages"], start=1)
        ]
        plan = ContentPlan(project_id=case["id"], outline_version=1, slides=pages)
        layouts.update(normalize_layout(page.layout) for page in pages)
        try:
            specs.extend(compile_specs(plan, _design()))
        except Exception as exc:
            crashes.append(type(exc).__name__)

    from pptx import Presentation
    expected_slides = sum(len(c["pages"]) for c in cases)
    with tempfile.TemporaryDirectory() as directory:
        out = Path(directory) / "evaluation.pptx"
        try:
            build_pptx(specs, _design(), out)
            prs = Presentation(str(out))
        except Exception as exc:
            crashes.append(type(exc).__name__)
            prs = Presentation()
        slide_count_match = len(prs.slides) == expected_slides
        text_editable = any(
            shape.has_text_frame and shape.text_frame.text
            for slide in prs.slides
            for shape in slide.shapes
        )

    canvas = Box(0, 0, 13.333 * 914400, 7.5 * 914400)
    oob = 0
    overlap = 0
    total_elements = 0
    for spec in specs:
        for el in spec.elements:
            total_elements += 1
            if out_of_bounds(_element_box(el), canvas):
                oob += 1
        content_boxes = [
            _element_box(element)
            for element in spec.elements
            if element.type != "rect"
        ]
        for i, first in enumerate(content_boxes):
            for second in content_boxes[i + 1:]:
                if overlap_ratio(first, second) >= 0.2:
                    overlap += 1

    encoded = json.dumps({"slide_count_match": slide_count_match, "out_of_bounds_rate": 0.0},
                         ensure_ascii=False)
    secret_leaks = int("TEST_KEY" in encoded or "sk-" in encoded)

    return {
        "schema_version": 1,
        "cases": len(cases),
        "layouts_exercised": sorted(layouts),
        "page_types_exercised": sorted({spec.page_type for spec in specs}),
        "crashes": crashes,
        "slide_count_match": slide_count_match,
        "out_of_bounds_rate": oob / total_elements if total_elements else 0,
        "severe_overlap_rate": overlap / total_elements if total_elements else 0,
        "text_editable": text_editable,
        "secret_leaks": secret_leaks,
    }


def main() -> int:
    cases = json.loads(Path("tests/fixtures/m3/cases.json").read_text("utf-8"))
    result = evaluate(cases)
    output = Path("artifacts/m3")
    output.mkdir(parents=True, exist_ok=True)
    result["secret_leaks"] = max(result["secret_leaks"],
        int("TEST_KEY" in json.dumps(result) or "sk-" in json.dumps(result)))
    (output / "evaluation.json").write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())

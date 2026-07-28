from pptx import Presentation

from pptmaker.evaluation.product_acceptance import evaluate_deck
from pptmaker.m3.blueprint_compiler import compile_blueprint, design_from_blueprint
from pptmaker.m3.builder import build_pptx
from pptmaker.m3.quality import inspect_specs
from tests.acceptance.fixtures import acceptance_inputs


def test_approved_blueprint_builds_reopenable_quality_deck(tmp_path):
    blueprint, facts, assets = acceptance_inputs(tmp_path)
    design = design_from_blueprint(blueprint)
    specs = compile_blueprint(blueprint, facts, assets, design)

    assert inspect_specs(specs).passed
    output = build_pptx(specs, design, tmp_path / "accepted.pptx")
    reopened = Presentation(str(output))
    report = evaluate_deck(blueprint, specs, output)

    assert len(reopened.slides) == len(blueprint.pages)
    assert any(shape.has_chart for shape in reopened.slides[6].shapes)
    assert not any(
        element.type == "image_placeholder"
        for spec in specs
        for element in spec.elements
    )
    assert report.passed

from datetime import datetime, timezone

import pytest
from PIL import Image
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from pptmaker.assets.models import AssetRecord
from pptmaker.blueprint.layout_planner import plan_page_layout
from pptmaker.blueprint.models import AssetSlot, LayoutRegion
from pptmaker.m3.blueprint_compiler import (
    compile_blueprint,
    design_from_blueprint,
)
from pptmaker.m3.builder import build_pptx
from tests.blueprint.factories import blueprint


def _slot(slot_id, asset_id):
    return AssetSlot(
        slot_id=slot_id,
        role="photo",
        subject=slot_id,
        search_query=slot_id,
        purpose="支撑页面结论",
        region=LayoutRegion(x=0.1, y=0.2, width=0.2, height=0.2),
        selected_asset_id=asset_id,
    )


def _asset(path, slot_id, asset_id, page_id="p1"):
    return AssetRecord(
        id=asset_id,
        project_id="p1",
        blueprint_version=1,
        page_id=page_id,
        slot_id=slot_id,
        provider="fixture",
        original_url=f"https://example.com/{asset_id}.png",
        landing_url=f"https://example.com/{asset_id}",
        thumbnail_url=f"https://example.com/{asset_id}-thumb.png",
        license_status="UNKNOWN",
        license_mode="LOOSE",
        width=640,
        height=360,
        accessed_at=datetime.now(timezone.utc),
        status="APPROVED",
        local_path=str(path),
        content_hash="a" * 64,
        mime_type="image/png",
        byte_size=path.stat().st_size,
    )


@pytest.mark.parametrize("count", [1, 2, 3, 4])
def test_multi_image_blueprint_builds_editable_pptx(tmp_path, count):
    deck = blueprint(page_count=1)
    slots = []
    assets = []
    for index in range(count):
        asset_id = f"asset-{index}"
        slot_id = f"slot-{index}"
        path = tmp_path / f"{asset_id}.png"
        Image.new(
            "RGB",
            (640, 360),
            (30 + index * 40, 80, 160),
        ).save(path)
        slots.append(_slot(slot_id, asset_id))
        assets.append(_asset(path, slot_id, asset_id))
    page = plan_page_layout(
        deck.pages[0].model_copy(
            update={"page_type": "image_text", "asset_slots": slots}
        )
    )
    approved = deck.model_copy(
        update={"status": "APPROVED", "pages": [page]}
    )
    design = design_from_blueprint(approved)
    specs = compile_blueprint(approved, [], assets, design)
    output = build_pptx(specs, design, tmp_path / "multi-image.pptx")

    reopened = Presentation(output)
    pictures = [
        shape
        for shape in reopened.slides[0].shapes
        if shape.shape_type == MSO_SHAPE_TYPE.PICTURE
    ]
    assert len(pictures) == count
    assert output.is_file()


def test_optional_missing_image_builds_without_placeholder(tmp_path):
    deck = blueprint(page_count=1)
    page = plan_page_layout(
        deck.pages[0].model_copy(
            update={"asset_slots": [_slot("optional", "missing")]}
        )
    )
    approved = deck.model_copy(
        update={"status": "APPROVED", "pages": [page]}
    )
    design = design_from_blueprint(approved)
    specs = compile_blueprint(approved, [], [], design)
    output = build_pptx(specs, design, tmp_path / "fallback.pptx")
    assert output.is_file()
    assert not any(
        element.type == "image_placeholder"
        for element in specs[0].elements
    )

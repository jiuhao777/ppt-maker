from pathlib import Path
from zipfile import ZipFile

import pytest
from PIL import Image
from pptx import Presentation
from pptx.util import Inches

from pptmaker.m3.models import (
    SlideSpec, TextElement, BulletElement, ImagePlaceholder, ImageElement,
    ChartElement, LineElement, RectShape, TableBlock,
    DesignPlan, ColorPalette, FontSet,
)
from pptmaker.m3.builder import build_pptx


def _design() -> DesignPlan:
    from pptmaker.m3.design import PRESETS
    return PRESETS["business"].model_copy(update={"project_id": "p1"})


def _spec(**kw) -> SlideSpec:
    defaults = dict(page_id="p1", order=1, page_type="content", background_color="#FFFFFF", elements=[])
    return SlideSpec(**(defaults | kw))


def _text(text="Hello", **kw) -> TextElement:
    defaults = dict(type="text", text=text, x=1.0, y=1.0, width=5.0, height=1.0,
                    font_size=16, font_name="Arial", color="#000000")
    return TextElement(**(defaults | kw))


class TestBuilder:
    def test_output_file_exists_and_non_empty(self, tmp_path):
        out = tmp_path / "out.pptx"
        result = build_pptx([_spec()], _design(), out)
        assert result == out
        assert out.stat().st_size > 0

    def test_slide_count_matches_input(self, tmp_path):
        out = tmp_path / "out.pptx"
        specs = [_spec(page_id=f"p{i}", order=i) for i in range(1, 4)]
        build_pptx(specs, _design(), out)
        prs = Presentation(str(out))
        assert len(prs.slides) == 3

    def test_pptx_is_16_9(self, tmp_path):
        out = tmp_path / "out.pptx"
        build_pptx([_spec()], _design(), out)
        prs = Presentation(str(out))
        assert prs.slide_width == Inches(13.333)
        assert prs.slide_height == Inches(7.5)

    def test_text_is_editable(self, tmp_path):
        out = tmp_path / "out.pptx"
        build_pptx([_spec(elements=[_text("Editable")])], _design(), out)
        prs = Presentation(str(out))
        shapes = [s for s in prs.slides[0].shapes if s.has_text_frame]
        assert len(shapes) >= 1
        assert shapes[0].text_frame.text == "Editable"

    def test_bullet_element_renders(self, tmp_path):
        out = tmp_path / "out.pptx"
        el = BulletElement(type="bullet", items=["A", "B"], x=1.0, y=1.0, width=5.0, height=2.0,
                           font_size=16, font_name="Arial", color="#000000")
        build_pptx([_spec(elements=[el])], _design(), out)
        prs = Presentation(str(out))
        text = prs.slides[0].shapes[0].text_frame.text
        assert "A" in text and "B" in text

    def test_rect_shape_renders(self, tmp_path):
        out = tmp_path / "out.pptx"
        el = RectShape(type="rect", x=1.0, y=1.0, width=2.0, height=1.0, fill_color="#FF0000")
        build_pptx([_spec(elements=[el])], _design(), out)
        prs = Presentation(str(out))
        assert len(prs.slides[0].shapes) >= 1

    def test_image_placeholder_is_rejected(self, tmp_path):
        out = tmp_path / "out.pptx"
        el = ImagePlaceholder(type="image_placeholder", x=1.0, y=1.0, width=3.0, height=2.0,
                              description="chart")
        with pytest.raises(ValueError, match="placeholders"):
            build_pptx([_spec(elements=[el])], _design(), out)
        assert not out.exists()

    def test_table_renders(self, tmp_path):
        out = tmp_path / "out.pptx"
        el = TableBlock(type="table", rows=[["H1", "H2"], ["D1", "D2"]],
                        x=1.0, y=1.0, width=5.0, height=2.0, font_size=12)
        build_pptx([_spec(elements=[el])], _design(), out)
        prs = Presentation(str(out))
        tables = [s for s in prs.slides[0].shapes if s.has_table]
        assert len(tables) == 1

    def test_verified_image_is_cropped_not_stretched(self, tmp_path):
        image_path = tmp_path / "wide.png"
        Image.new("RGB", (1200, 400), "blue").save(image_path)
        element = ImageElement(
            path=str(image_path), x=1, y=1, width=4, height=3
        )
        out = build_pptx(
            [_spec(elements=[element])], _design(), tmp_path / "out.pptx"
        )
        picture = Presentation(str(out)).slides[0].shapes[0]
        assert picture.crop_left > 0
        assert picture.crop_left == pytest.approx(picture.crop_right)

    def test_native_chart_remains_editable(self, tmp_path):
        element = ChartElement(
            chart_type="column",
            categories=["A", "B"],
            series={"Value": [1, 2]},
            x=1,
            y=1,
            width=6,
            height=4,
        )
        out = build_pptx(
            [_spec(elements=[element])], _design(), tmp_path / "out.pptx"
        )
        assert Presentation(str(out)).slides[0].shapes[0].has_chart

    def test_notes_sources_footer_line_and_transition(self, tmp_path):
        line = LineElement(
            x=1, y=2, width=4, height=0, color="#FF0000"
        )
        spec = _spec(
            elements=[line],
            speaker_notes="Talk track",
            source_references=["Report · https://example.com"],
            transition="fade",
        )
        out = build_pptx([spec], _design(), tmp_path / "out.pptx")
        reopened = Presentation(str(out))
        notes = reopened.slides[0].notes_slide.notes_text_frame.text
        assert "Talk track" in notes
        assert "[Sources]" in notes
        assert "来源：" in reopened.slides[0].shapes[-1].text
        with ZipFile(out) as archive:
            xml = archive.read("ppt/slides/slide1.xml").decode("utf-8")
        assert "<p:transition>" in xml and "<p:fade" in xml

    def test_missing_image_leaves_no_published_file(self, tmp_path):
        out = tmp_path / "out.pptx"
        element = ImageElement(
            path=str(tmp_path / "missing.png"),
            x=1,
            y=1,
            width=4,
            height=3,
        )
        with pytest.raises(ValueError, match="missing"):
            build_pptx([_spec(elements=[element])], _design(), out)
        assert not out.exists()

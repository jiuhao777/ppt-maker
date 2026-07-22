from io import BytesIO
from pathlib import Path

from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.chart.data import ChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from pptx.util import Inches, Pt

from pptmaker.m0.contracts import ProbeResult, ProbeStatus


def _build_probe_image() -> BytesIO:
    image = Image.new("RGB", (640, 360), "#EAF2F8")
    draw = ImageDraw.Draw(image)
    draw.rectangle((48, 48, 592, 312), fill="#1F4E79")
    draw.ellipse((240, 100, 400, 260), fill="#5B9BD5")
    stream = BytesIO()
    image.save(stream, format="PNG")
    stream.seek(0)
    return stream


def build_baseline_deck(output: Path) -> ProbeResult:
    presentation = Presentation()
    presentation.slide_width = Inches(13.333)
    presentation.slide_height = Inches(7.5)

    title_slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    title = title_slide.shapes.add_textbox(
        Inches(0.8), Inches(0.7), Inches(11.7), Inches(0.8)
    )
    title_run = title.text_frame.paragraphs[0].add_run()
    title_run.text = "PPT Maker 技术验证"
    title_run.font.size = Pt(30)
    title_run.font.bold = True

    card = title_slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE,
        Inches(0.8),
        Inches(2.0),
        Inches(4.0),
        Inches(2.0),
    )
    card.fill.solid()
    card.fill.fore_color.rgb = RGBColor(31, 78, 121)
    card.text = "文本与形状保持可编辑"
    title_slide.shapes.add_picture(
        _build_probe_image(), Inches(6.0), Inches(2.0), width=Inches(5.4)
    )

    data_slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    heading = data_slide.shapes.add_textbox(
        Inches(0.7), Inches(0.25), Inches(11.8), Inches(0.5)
    )
    heading.text_frame.text = "原生表格与图表"

    table = data_slide.shapes.add_table(
        3, 2, Inches(0.7), Inches(0.9), Inches(4.2), Inches(2.2)
    ).table
    for row, values in enumerate((("阶段", "得分"), ("生成", "80"), ("修复", "92"))):
        for column, value in enumerate(values):
            table.cell(row, column).text = value

    chart_data = ChartData()
    chart_data.categories = ["生成", "修复"]
    chart_data.add_series("质量", (80, 92))
    data_slide.shapes.add_chart(
        XL_CHART_TYPE.COLUMN_CLUSTERED,
        Inches(5.3),
        Inches(0.9),
        Inches(7.2),
        Inches(4.8),
        chart_data,
    )

    output.parent.mkdir(parents=True, exist_ok=True)
    presentation.save(output)
    reopened = Presentation(output)
    if len(reopened.slides) != 2:
        raise RuntimeError("saved presentation did not reopen with two slides")

    return ProbeResult(
        name="pptx-baseline",
        status=ProbeStatus.PASS,
        metrics={
            "slides": 2,
            "text_boxes": 2,
            "shapes": 1,
            "tables": 1,
            "images": 1,
            "charts": 1,
        },
        evidence=[str(output.resolve())],
    )

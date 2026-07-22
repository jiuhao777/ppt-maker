import time
from pathlib import Path

from pydantic import BaseModel, Field

from pptmaker.m0.contracts import ProbeResult, ProbeStatus


SENTINEL = "M0_SOURCE_SENTINEL_20260722"


class ParsedDocument(BaseModel):
    kind: str
    text: str
    tables: list[list[str]] = Field(default_factory=list)
    images: int = 0
    locations: list[str] = Field(default_factory=list)


def parse_markdown(content: str) -> ParsedDocument:
    text_lines: list[str] = []
    table_rows: list[list[str]] = []
    locations: list[str] = []
    in_code_fence = False
    for line_number, raw_line in enumerate(content.splitlines(), start=1):
        line = raw_line.strip()
        if not line:
            continue
        if line.startswith("```"):
            in_code_fence = not in_code_fence
            text_lines.append(line)
            locations.append(f"line:{line_number}")
            continue
        if line.startswith("|") and not in_code_fence:
            cells = [cell.strip() for cell in line.strip("|").split("|")]
            if cells and not all(set(cell) <= {"-", ":"} for cell in cells):
                table_rows.append(cells)
                locations.append(f"line:{line_number}")
            continue
        text_lines.append(line.lstrip("# "))
        locations.append(f"line:{line_number}")
    return ParsedDocument(
        kind="markdown",
        text="\n".join(text_lines),
        tables=table_rows,
        locations=locations,
    )


def parse_pptx(path: Path) -> ParsedDocument:
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    texts: list[str] = []
    tables: list[list[str]] = []
    locations: list[str] = []
    image_count = 0

    def visit_shapes(shapes, prefix: str) -> None:
        nonlocal image_count
        for shape_index, shape in enumerate(shapes, start=1):
            location = f"{prefix}/shape:{shape_index}"
            if getattr(shape, "shape_type", None) == MSO_SHAPE_TYPE.PICTURE:
                image_count += 1
            if getattr(shape, "has_text_frame", False) and shape.text.strip():
                texts.append(shape.text.strip())
                locations.append(location)
            if getattr(shape, "has_table", False):
                for row_index, row in enumerate(shape.table.rows, start=1):
                    tables.append([cell.text.strip() for cell in row.cells])
                    locations.append(f"{location}/table-row:{row_index}")
            if getattr(shape, "shape_type", None) == MSO_SHAPE_TYPE.GROUP:
                visit_shapes(shape.shapes, f"{prefix}/group:{shape_index}")

    presentation = Presentation(path)
    for slide_index, slide in enumerate(presentation.slides, start=1):
        visit_shapes(slide.shapes, f"slide:{slide_index}")
        if slide.has_notes_slide:
            for shape_index, shape in enumerate(slide.notes_slide.shapes, start=1):
                if getattr(shape, "has_text_frame", False) and shape.text.strip():
                    text = shape.text.strip()
                    if text not in texts:
                        texts.append(text)
                        locations.append(f"slide:{slide_index}/notes:shape:{shape_index}")
    return ParsedDocument(
        kind="pptx",
        text="\n".join(texts),
        tables=tables,
        images=image_count,
        locations=locations,
    )


def parse_with_markitdown(path: Path) -> ParsedDocument:
    from markitdown import MarkItDown

    result = MarkItDown().convert_local(path)
    text = getattr(result, "markdown", None) or getattr(result, "text_content", "")
    return ParsedDocument(kind=path.suffix.lstrip(".").lower(), text=text or "")


def generate_parser_fixtures(root: Path) -> None:
    from docx import Document
    from PIL import Image, ImageDraw
    from pptx import Presentation
    from pptx.util import Inches
    from reportlab.pdfgen.canvas import Canvas

    root.mkdir(parents=True, exist_ok=True)
    (root / "sample.md").write_text(f"# M0\n\n{SENTINEL}\n", encoding="utf-8")

    image = Image.new("RGB", (640, 360), "white")
    ImageDraw.Draw(image).text((20, 160), SENTINEL, fill="black")
    image.save(root / "sample.png")

    document = Document()
    document.add_heading("M0", 1)
    document.add_paragraph(SENTINEL)
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = "指标", "值"
    table.cell(1, 0).text, table.cell(1, 1).text = "质量", "90"
    document.save(root / "sample.docx")

    pdf = Canvas(str(root / "sample.pdf"))
    pdf.drawString(72, 760, SENTINEL)
    pdf.showPage()
    pdf.save()

    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    slide.shapes.add_textbox(Inches(1), Inches(1), Inches(8), Inches(1)).text = SENTINEL
    presentation.save(root / "sample.pptx")

    challenge = Presentation()
    challenge_slide = challenge.slides.add_slide(challenge.slide_layouts[6])
    group = challenge_slide.shapes.add_group_shape()
    group.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1)).text = "GROUP_SENTINEL"
    table_shape = challenge_slide.shapes.add_table(
        2, 2, Inches(1), Inches(3), Inches(5), Inches(2)
    )
    table_shape.table.cell(0, 0).text, table_shape.table.cell(0, 1).text = "指标", "值"
    table_shape.table.cell(1, 0).text, table_shape.table.cell(1, 1).text = "质量", "90"
    challenge_slide.notes_slide.notes_text_frame.text = "NOTES_SENTINEL"
    challenge.save(root / "challenge.pptx")


def run_parser_core_probe(root: Path) -> ProbeResult:
    fixture_root = root / "artifacts" / "m0" / "parser-fixtures"
    generate_parser_fixtures(fixture_root)
    started = time.perf_counter()
    markdown = parse_markdown((fixture_root / "sample.md").read_text("utf-8"))
    powerpoint = parse_pptx(fixture_root / "sample.pptx")
    elapsed = time.perf_counter() - started
    checks = {
        "markdown_text": SENTINEL in markdown.text,
        "pptx_text": SENTINEL in powerpoint.text,
        "pptx_location": bool(powerpoint.locations),
    }
    return ProbeResult(
        name="parser-core",
        status=ProbeStatus.PASS if all(checks.values()) else ProbeStatus.FAIL,
        metrics={
            "required_checks": len(checks),
            "passed_checks": sum(checks.values()),
            "elapsed_seconds": elapsed,
        },
        details={"checks": checks},
    )


def measure_parser_matrix(root: Path) -> list[dict[str, str | bool | float]]:
    fixture_root = root / "artifacts" / "m0" / "parser-matrix-fixtures"
    generate_parser_fixtures(fixture_root)
    rows: list[dict[str, str | bool | float]] = []

    def measure(
        name: str,
        dependency: str,
        operation,
        expected: str,
        require_location: bool,
    ) -> None:
        started = time.perf_counter()
        try:
            parsed = operation()
            normalized = parsed.text.replace("\\_", "_")
            text_found = expected in normalized
            location_found = bool(parsed.locations)
            if text_found and (location_found or not require_location):
                status = "pass"
            elif text_found:
                status = "partial"
            else:
                status = "unsupported"
            error = ""
        except Exception as exc:
            text_found = location_found = False
            status = "failed"
            error = f"{type(exc).__name__}: {exc}"
        rows.append(
            {
                "format": name,
                "dependency": dependency,
                "status": status,
                "text_found": text_found,
                "location_found": location_found,
                "elapsed_seconds": time.perf_counter() - started,
                "error": error,
            }
        )

    measure(
        "markdown-core",
        "deterministic",
        lambda: parse_markdown((fixture_root / "sample.md").read_text("utf-8")),
        SENTINEL,
        True,
    )
    measure(
        "pptx-core-challenge",
        "python-pptx",
        lambda: parse_pptx(fixture_root / "challenge.pptx"),
        "NOTES_SENTINEL",
        True,
    )
    for name, expected in (
        ("sample.docx", SENTINEL),
        ("sample.pdf", SENTINEL),
        ("sample.pptx", SENTINEL),
        ("challenge.pptx", "NOTES_SENTINEL"),
        ("sample.png", SENTINEL),
    ):
        measure(
            f"markitdown-{Path(name).suffix.lstrip('.')}-{Path(name).stem}",
            "markitdown",
            lambda name=name: parse_with_markitdown(fixture_root / name),
            expected,
            True,
        )
    return rows

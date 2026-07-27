import os
import tempfile
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.chart.data import ChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE, MSO_CONNECTOR
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt

from pptmaker.m3.image_fit import crop_to_fill
from pptmaker.m3.models import DesignPlan, SlideSpec


_CHART_TYPES = {
    "column": XL_CHART_TYPE.COLUMN_CLUSTERED,
    "bar": XL_CHART_TYPE.BAR_CLUSTERED,
    "line": XL_CHART_TYPE.LINE_MARKERS,
    "pie": XL_CHART_TYPE.PIE,
    "doughnut": XL_CHART_TYPE.DOUGHNUT,
}


def build_pptx(
    specs: list[SlideSpec],
    design: DesignPlan,
    output_path: Path,
) -> Path:
    if any(
        element.type == "image_placeholder"
        for spec in specs
        for element in spec.elements
    ):
        raise ValueError("image placeholders cannot be exported")
    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    descriptor, temporary_name = tempfile.mkstemp(
        prefix=f".{output_path.stem}-",
        suffix=".tmp.pptx",
        dir=output_path.parent,
    )
    os.close(descriptor)
    temporary = Path(temporary_name)
    try:
        presentation = Presentation()
        presentation.slide_width = Inches(13.333)
        presentation.slide_height = Inches(7.5)
        blank = presentation.slide_layouts[6]
        for spec in specs:
            slide = presentation.slides.add_slide(blank)
            _set_bg(slide, spec.background_color)
            for element in spec.elements:
                _render_element(slide, element)
            _add_sources(slide, spec.source_references, design)
            _add_notes(
                slide, spec.speaker_notes, spec.source_references
            )
            if spec.transition == "fade":
                _add_fade(slide)
        presentation.save(str(temporary))
        reopened = Presentation(str(temporary))
        if len(reopened.slides) != len(specs):
            raise ValueError("saved slide count does not match")
        os.replace(temporary, output_path)
    except Exception:
        temporary.unlink(missing_ok=True)
        raise
    return output_path


def _render_element(slide, element):
    if element.type == "text":
        _add_text(slide, element)
    elif element.type == "bullet":
        _add_bullets(slide, element)
    elif element.type == "rect":
        _add_rect(slide, element)
    elif element.type == "table":
        _add_table(slide, element)
    elif element.type == "image":
        _add_image(slide, element)
    elif element.type == "chart":
        _add_chart(slide, element)
    elif element.type == "line":
        _add_line(slide, element)
    else:
        raise ValueError(f"unsupported slide element: {element.type}")


def _set_bg(slide, color: str):
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = _rgb(color)


def _add_text(slide, element):
    box = slide.shapes.add_textbox(
        Inches(element.x),
        Inches(element.y),
        Inches(element.width),
        Inches(element.height),
    )
    frame = box.text_frame
    frame.word_wrap = True
    paragraph = frame.paragraphs[0]
    paragraph.text = element.text
    paragraph.font.size = Pt(element.font_size)
    paragraph.font.bold = element.bold
    paragraph.font.color.rgb = _rgb(element.color)
    paragraph.font.name = element.font_name


def _add_bullets(slide, element):
    box = slide.shapes.add_textbox(
        Inches(element.x),
        Inches(element.y),
        Inches(element.width),
        Inches(element.height),
    )
    frame = box.text_frame
    frame.word_wrap = True
    for index, item in enumerate(element.items):
        paragraph = (
            frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        )
        paragraph.text = f"• {item}"
        paragraph.font.size = Pt(element.font_size)
        paragraph.font.color.rgb = _rgb(element.color)
        paragraph.font.name = element.font_name
        paragraph.space_after = Pt(10)


def _add_rect(slide, element):
    shape = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE,
        Inches(element.x),
        Inches(element.y),
        Inches(element.width),
        Inches(element.height),
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = _rgb(element.fill_color)
    shape.line.fill.background()


def _add_table(slide, element):
    column_count = len(element.rows[0])
    if any(len(row) != column_count for row in element.rows):
        raise ValueError("table rows must have equal length")
    shape = slide.shapes.add_table(
        len(element.rows),
        column_count,
        Inches(element.x),
        Inches(element.y),
        Inches(element.width),
        Inches(element.height),
    )
    for row_index, row in enumerate(element.rows):
        for column_index, text in enumerate(row):
            cell = shape.table.cell(row_index, column_index)
            cell.text = text
            for paragraph in cell.text_frame.paragraphs:
                paragraph.font.size = Pt(element.font_size)


def _add_image(slide, element):
    path = Path(element.path)
    if not path.is_file():
        raise ValueError("image file is missing")
    with Image.open(path) as image:
        source_width, source_height = image.size
    left, top, right, bottom = crop_to_fill(
        source_width,
        source_height,
        element.width,
        element.height,
    )
    picture = slide.shapes.add_picture(
        str(path),
        Inches(element.x),
        Inches(element.y),
        Inches(element.width),
        Inches(element.height),
    )
    picture.crop_left = left
    picture.crop_top = top
    picture.crop_right = right
    picture.crop_bottom = bottom


def _add_chart(slide, element):
    if any(
        len(values) != len(element.categories)
        for values in element.series.values()
    ):
        raise ValueError("chart series must match categories")
    data = ChartData()
    data.categories = element.categories
    for name, values in element.series.items():
        data.add_series(name, values)
    slide.shapes.add_chart(
        _CHART_TYPES[element.chart_type],
        Inches(element.x),
        Inches(element.y),
        Inches(element.width),
        Inches(element.height),
        data,
    )


def _add_line(slide, element):
    line = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT,
        Inches(element.x),
        Inches(element.y),
        Inches(element.x + element.width),
        Inches(element.y + element.height),
    )
    line.line.color.rgb = _rgb(element.color)
    line.line.width = Pt(element.line_width)


def _add_sources(slide, sources: list[str], design: DesignPlan):
    if not sources:
        return
    text = "来源：" + "；".join(sources)
    box = slide.shapes.add_textbox(
        Inches(0.7), Inches(7.15), Inches(11.9), Inches(0.2)
    )
    paragraph = box.text_frame.paragraphs[0]
    paragraph.text = text[:500]
    paragraph.font.size = Pt(9)
    paragraph.font.name = design.fonts.body_zh
    paragraph.font.color.rgb = _rgb(design.palette.text_secondary)


def _add_notes(slide, notes: str, sources: list[str]):
    blocks = []
    if notes:
        blocks.append(notes)
    if sources:
        blocks.append("[Sources]\n" + "\n".join(sources))
    if blocks:
        slide.notes_slide.notes_text_frame.text = "\n\n".join(blocks)


def _add_fade(slide):
    transition = OxmlElement("p:transition")
    transition.append(OxmlElement("p:fade"))
    slide._element.insert_element_before(
        transition, "p:timing", "p:extLst"
    )


def _rgb(hex_color: str) -> RGBColor:
    value = hex_color.lstrip("#")
    return RGBColor(
        int(value[0:2], 16),
        int(value[2:4], 16),
        int(value[4:6], 16),
    )
